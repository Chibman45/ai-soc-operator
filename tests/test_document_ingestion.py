"""Security and extraction contracts for uploaded organization documents."""

from __future__ import annotations

import base64
import io
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from web.app import app, hash_password, init_db
from web.services.document_ingestion import (
    SUPPORTED_EXTENSIONS,
    DocumentValidationError,
    extract_document,
    ingest_document,
    mark_document_analyzed,
)
from web.services.schema_features import init_feature_db

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUB"
    "AScY42YAAAAASUVORK5CYII="
)


def ingestion_connection() -> tuple[sqlite3.Connection, int]:
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript(
        """
        CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT UNIQUE, password_hash TEXT);
        CREATE TABLE credentials (id INTEGER PRIMARY KEY, platform TEXT, key_name TEXT, key_value TEXT);
        CREATE TABLE platform_config (id INTEGER PRIMARY KEY, platform TEXT UNIQUE, base_url TEXT, enabled INTEGER, updated_at TIMESTAMP);
        CREATE TABLE runs (id INTEGER PRIMARY KEY);
        CREATE TABLE cases (id INTEGER PRIMARY KEY, title TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE activity_log (id INTEGER PRIMARY KEY, run_id INTEGER);
        INSERT INTO users VALUES (1, 'admin', 'x');
        """
    )
    init_feature_db(conn)
    org_id = int(conn.execute("SELECT id FROM organizations ORDER BY id LIMIT 1").fetchone()[0])
    return conn, org_id


class DocumentValidationTests(unittest.TestCase):
    def test_supported_formats_match_the_approved_plan(self):
        self.assertEqual(
            SUPPORTED_EXTENSIONS,
            {
                ".yaml", ".yml", ".json", ".pdf", ".docx", ".md", ".txt",
                ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp",
            },
        )

    def test_rejects_extension_content_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            fake = Path(tmp) / "architecture.pdf"
            fake.write_text("This is not a PDF", encoding="utf-8")
            with self.assertRaisesRegex(DocumentValidationError, "signature"):
                extract_document(fake, original_name=fake.name)

    def test_rejects_path_like_original_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "safe.txt"
            source.write_text("safe", encoding="utf-8")
            with self.assertRaisesRegex(DocumentValidationError, "file name"):
                extract_document(source, original_name="../unsafe.txt")

    def test_rejects_declared_mime_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "architecture.pdf"
            source.write_bytes(b"%PDF-1.7\nmock")
            with self.assertRaisesRegex(DocumentValidationError, "MIME"):
                extract_document(source, original_name=source.name, declared_mime="image/png")

    def test_malformed_pdf_fails_as_a_bounded_runtime_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "malformed.pdf"
            source.write_bytes(b"%PDF-1.7\nnot a real object graph")
            with self.assertRaisesRegex(RuntimeError, "PDF native text extraction failed"):
                extract_document(source, original_name=source.name)

    def test_rejects_oversized_extracted_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "oversized.txt"
            source.write_text("x" * 500_001, encoding="utf-8")
            with self.assertRaisesRegex(DocumentValidationError, "character limit"):
                extract_document(source, original_name=source.name)


class ExtractionTests(unittest.TestCase):
    def test_text_extraction_includes_provenance(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "policy.md"
            source.write_text("# Access policy\nAnalysts are read only.", encoding="utf-8")
            result = extract_document(source, original_name=source.name)
        self.assertIn("Analysts are read only", result["text"])
        self.assertEqual(result["provenance"][0]["kind"], "text")
        self.assertFalse(result["ocr_used"])

    def test_docx_includes_body_table_header_footer_and_embedded_image_ocr(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "chart.png"
            image.write_bytes(PNG_1X1)
            source = root / "organization.docx"
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "Security Operations"
            doc.add_paragraph("Alice leads the SOC")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Asset"
            table.cell(0, 1).text = "Owner"
            doc.add_picture(str(image))
            doc.sections[0].footer.paragraphs[0].text = "Confidential"
            doc.save(source)

            with patch(
                "web.services.document_ingestion._ocr_image",
                return_value={
                    "text": "SOC reports to CISO",
                    "boxes": [{"text": "CISO", "left": 1, "top": 1, "width": 2, "height": 2}],
                },
            ):
                result = extract_document(source, original_name=source.name)

        self.assertIn("Security Operations", result["text"])
        self.assertIn("Alice leads the SOC", result["text"])
        self.assertIn("Asset | Owner", result["text"])
        self.assertIn("Confidential", result["text"])
        self.assertIn("SOC reports to CISO", result["text"])
        self.assertTrue(result["ocr_used"])
        self.assertIn("image", {item["kind"] for item in result["provenance"]})

    def test_pdf_uses_ocr_only_for_pages_without_native_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "network.pdf"
            source.write_bytes(b"%PDF-1.7\nmock")
            with (
                patch(
                    "web.services.document_ingestion._extract_pdf_pages",
                    return_value=["Native page text is sufficiently descriptive.", ""],
                ),
                patch("web.services.document_ingestion._render_pdf_page", return_value=object()) as render,
                patch(
                    "web.services.document_ingestion._ocr_image",
                    return_value={"text": "DMZ to SIEM", "boxes": []},
                ),
            ):
                result = extract_document(source, original_name=source.name)

        render.assert_called_once_with(source, 2)
        self.assertIn("Native page text", result["text"])
        self.assertIn("DMZ to SIEM", result["text"])
        self.assertEqual([item["method"] for item in result["provenance"]], ["native", "ocr"])


class IngestionLifecycleTests(unittest.TestCase):
    def test_ingestion_copies_immutable_evidence_and_requires_preview_before_analysis(self):
        conn, org_id = ingestion_connection()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "policy.txt"
            source.write_text("Human approval is required.", encoding="utf-8")
            evidence = root / "evidence"
            record = ingest_document(
                conn,
                org_id=org_id,
                purpose="organization_policy",
                source_path=source,
                source_name=source.name,
                created_by=1,
                evidence_root=evidence,
            )
            stored = Path(record["evidence_path"])
            self.assertTrue(stored.is_file())
            self.assertEqual(record["status"], "preview")
            self.assertEqual(stored.read_text(encoding="utf-8"), "Human approval is required.")

            source.write_text("Changed after upload", encoding="utf-8")
            self.assertEqual(stored.read_text(encoding="utf-8"), "Human approval is required.")
            analyzed = mark_document_analyzed(conn, record["id"], org_id=org_id, result_type="policy_draft", result_id=7)

        self.assertEqual(analyzed["status"], "analyzed")
        with self.assertRaisesRegex(ValueError, "preview"):
            mark_document_analyzed(conn, record["id"], org_id=org_id, result_type="policy_draft", result_id=8)


class DocumentRouteTests(unittest.TestCase):
    def test_playbook_upload_requires_preview_analysis_and_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            database = root / "soc_operator.db"
            (root / "playbooks").mkdir()
            with (
                patch("web.app.DB_PATH", database),
                patch("web.security.DB_PATH", database),
                patch("web.routes.documents.DB_PATH", database),
                patch("web.routes.admin.DB_PATH", database),
                patch("web.routes.assurance.DB_PATH", database),
                patch("web.app.ROOT", root),
                patch("web.routes.documents.ROOT", root),
            ):
                init_db()
                conn = sqlite3.connect(database)
                owner_id = int(
                    conn.execute(
                        "INSERT INTO users (username, password_hash) VALUES ('owner', ?)",
                        (hash_password("owner-password-123"),),
                    ).lastrowid
                )
                conn.commit()
                conn.close()
                init_db()
                conn = sqlite3.connect(database)
                org_id = int(conn.execute("SELECT default_org_id FROM users WHERE id = ?", (owner_id,)).fetchone()[0])
                conn.close()

                client = app.test_client()
                with client.session_transaction() as sess:
                    sess["user_id"] = owner_id
                    sess["org_id"] = org_id
                    sess["session_version"] = 0
                    sess["csrf_token"] = "document-csrf"

                playbook = b"id: reviewed-upload\nname: Reviewed upload\nsteps:\n  - id: report\n    type: report\n"
                response = client.post(
                    "/playbooks/upload",
                    data={
                        "csrf_token": "document-csrf",
                        "playbook": (io.BytesIO(playbook), "reviewed-upload.yaml"),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 302)
                self.assertIn("/documents/", response.headers["Location"])
                self.assertFalse((root / "playbooks" / "reviewed-upload.yaml").exists())

                preview = client.get(response.headers["Location"])
                self.assertEqual(preview.status_code, 200)
                self.assertIn(b"AI analysis has not run yet", preview.data)
                ingestion_id = int(response.headers["Location"].split("/")[-2])

                analyzed = client.post(
                    f"/api/documents/{ingestion_id}/analyze",
                    data={"csrf_token": "document-csrf"},
                )
                self.assertEqual(analyzed.status_code, 302, analyzed.get_data(as_text=True))
                conn = sqlite3.connect(database)
                draft_id, status = conn.execute(
                    "SELECT id, status FROM playbook_drafts WHERE ingestion_id = ?", (ingestion_id,)
                ).fetchone()
                conn.close()
                self.assertEqual(status, "draft")
                self.assertFalse((root / "playbooks" / "reviewed-upload.yaml").exists())

                approved = client.post(
                    f"/api/playbook-drafts/{draft_id}/approve",
                    headers={"X-CSRF-Token": "document-csrf"},
                )
                self.assertEqual(approved.status_code, 200, approved.get_data(as_text=True))
                self.assertTrue((root / "playbooks" / "reviewed-upload.yaml").is_file())

                policy_upload = client.post(
                    "/admin/policy/upload",
                    data={
                        "csrf_token": "document-csrf",
                        "policy": (io.BytesIO(b"Analysts may review alerts."), "organization.txt"),
                    },
                    content_type="multipart/form-data",
                )
                architecture_upload = client.post(
                    "/admin/architecture/upload",
                    data={
                        "csrf_token": "document-csrf",
                        "document": (io.BytesIO(b"DMZ forwards logs to Wazuh."), "architecture.txt"),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(policy_upload.status_code, 302)
                self.assertEqual(architecture_upload.status_code, 302)
                conn = sqlite3.connect(database)
                purposes = {
                    row[0]: row[1]
                    for row in conn.execute(
                        "SELECT purpose, status FROM document_ingestions WHERE purpose != 'playbook'"
                    ).fetchall()
                }
                self.assertEqual(
                    conn.execute("SELECT COUNT(*) FROM organization_policy_drafts").fetchone()[0], 0
                )
                self.assertEqual(
                    conn.execute("SELECT COUNT(*) FROM organization_architecture_drafts").fetchone()[0], 0
                )
                conn.close()
                self.assertEqual(purposes, {"organization_policy": "preview", "architecture": "preview"})


if __name__ == "__main__":
    raise SystemExit(unittest.main())
