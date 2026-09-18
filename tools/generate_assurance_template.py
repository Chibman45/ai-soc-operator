"""Generate the downloadable organization assurance input template."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "templates" / "organization-architecture-and-detection-assurance-template.docx"
NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F2F7FB"
GRAY = "D9D9D9"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    element = tbl_pr.find(qn("w:tblBorders"))
    if element is None:
        element = OxmlElement("w:tblBorders")
        tbl_pr.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), GRAY)
        element.append(border)


def margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    keep_with_next(paragraph)


def intro(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.style = doc.styles["Body Text"]


def form_table(doc: Document, rows: list[tuple[str, str]], widths=(2.25, 4.45)) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    for index, (label, guidance) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        cells[0].text = label
        cells[1].text = guidance
        shade(cells[0], LIGHT_BLUE)
        if index % 2:
            shade(cells[1], PALE_BLUE)
        for cell in cells:
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph()


def grid_table(doc: Document, headers: list[str], sample_rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.text = header
        shade(cell, NAVY)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)
    for row_index, values in enumerate(sample_rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            cells[index].text = value
            margins(cells[index], top=90, bottom=90)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[index], PALE_BLUE)
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    borders(table)
    doc.add_paragraph()


def checklist(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run("☐ ").font.name = "Arial"
        paragraph.add_run(item)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Body Text"].font.name = "Aptos"
    styles["Body Text"].font.size = Pt(10)
    styles["Body Text"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Body Text"].paragraph_format.space_after = Pt(8)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(27)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph(style="Title")
    title.add_run("Organization Architecture and Detection Assurance Input Template")
    subtitle = doc.add_paragraph()
    subtitle.add_run("AI SOC Operator").bold = True
    subtitle.add_run("  |  Controlled input for architecture understanding, detection validation, and platform assessment")
    subtitle.paragraph_format.space_after = Pt(18)

    intro(
        doc,
        "Complete this document to give AI SOC Operator a reviewed description of the organization, its assets, expected behavior, telemetry, detection rules, and SOC platform configuration. The system treats uploaded documents as untrusted input and creates a draft for administrator review. Uploading this document does not change access controls, rules, integrations, skills, or responders.",
    )
    warning = doc.add_paragraph()
    warning.add_run("Security requirement  ").bold = True
    warning.add_run("Do not enter passwords, API keys, private keys, authentication cookies, recovery codes, or other secrets. Store credentials only through the protected Credentials page.")
    warning.paragraph_format.space_after = Pt(14)

    heading(doc, "Document Control", 1)
    form_table(
        doc,
        [
            ("Organization name", "Enter the legal or operating name."),
            ("Document owner", "Name, role, and team responsible for accuracy."),
            ("Version", "Example: 1.0"),
            ("Prepared date", "Use YYYY-MM-DD."),
            ("Review date", "Use YYYY-MM-DD."),
            ("Classification", "Example: Internal, Confidential, or Restricted."),
            ("Approvers", "Name and role of architecture, SOC, and security approvers."),
        ],
    )

    heading(doc, "Completion and Approval", 1)
    checklist(
        doc,
        [
            "Asset owners verified the systems and business criticality in this document.",
            "Network and identity owners verified trust zones, domains, and expected data flows.",
            "SOC engineers verified log sources, rule identifiers, integrations, and retention settings.",
            "No credentials or other authentication secrets are included.",
            "An administrator will review the extracted architecture draft before activation.",
        ],
    )

    doc.add_page_break()
    heading(doc, "Organization Overview", 1)
    form_table(
        doc,
        [
            ("Operating model", "Describe business functions, locations, working hours, and support coverage."),
            ("SOC model", "Internal, outsourced, hybrid, follow-the-sun, or other model."),
            ("SOC hours and timezone", "State normal coverage and escalation availability."),
            ("Regulatory obligations", "List applicable standards, laws, contractual duties, and reporting deadlines."),
            ("Critical business services", "List services whose loss or compromise creates material impact."),
            ("Incident severity model", "Describe severity levels and who may declare each level."),
        ],
    )

    heading(doc, "Asset Inventory", 1)
    intro(doc, "Give every asset a stable identifier. Include servers, endpoints, network devices, identity systems, cloud services, security appliances, and business applications that affect detection decisions.")
    grid_table(
        doc,
        ["Asset ID", "Name and type", "Criticality", "Owner", "Trust zone", "Addresses", "OS or service", "Log sources"],
        [
            ["Example dc-01", "Domain Controller\nServer", "Critical", "Identity Team", "Server", "10.0.0.10", "Windows Server", "Wazuh, Windows Event"],
            ["", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", ""],
        ],
        [0.75, 1.0, 0.62, 0.75, 0.7, 0.8, 0.9, 1.0],
    )

    heading(doc, "Networks and Trust Zones", 1)
    grid_table(
        doc,
        ["Zone", "CIDRs", "Purpose", "Ingress controls", "Egress controls", "Logging", "Owner"],
        [["Example server", "10.0.0.0/24", "Internal services", "Firewall allowlist", "Proxy and firewall", "Firewall, DNS, Wazuh", "Infrastructure"], ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""]],
        [0.8, 0.9, 1.0, 1.05, 1.05, 1.05, 0.75],
    )

    heading(doc, "Identity Architecture", 1)
    form_table(
        doc,
        [
            ("Identity domains", "List Active Directory domains, Entra tenants, LDAP domains, and cloud identity tenants."),
            ("Privileged identity", "Describe administrator tiers, PAM, break-glass accounts, and approval requirements."),
            ("Authentication", "Describe MFA, conditional access, federation, SSO, and passwordless controls."),
            ("Service accounts", "List expected service-account patterns, owners, source systems, and normal schedules."),
            ("Joiner mover leaver", "Describe provisioning, role change, suspension, and termination workflows."),
        ],
    )

    heading(doc, "Expected Data Flows", 1)
    intro(doc, "Record expected communications so the system can distinguish normal behavior from suspicious activity. Do not list a flow as expected solely because it is currently observed.")
    grid_table(
        doc,
        ["Source asset or zone", "Destination", "Protocols and ports", "Purpose", "Schedule", "Expected", "Owner"],
        [["Example endpoints", "dc-01", "TCP 88, 389, 445", "Identity and policy", "Always", "Yes", "Identity Team"], ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""]],
        [1.0, 1.0, 1.05, 1.1, 0.75, 0.55, 0.75],
    )

    heading(doc, "Expected Behavior and Approved Exceptions", 1)
    grid_table(
        doc,
        ["Behavior ID", "Users or assets", "Expected behavior", "Schedule", "Approval or ticket", "Expiry", "Monitoring requirement"],
        [["Example EB-001", "backup-svc", "Large encrypted transfers to backup vault", "Daily 01:00", "CHG-0001", "Annual review", "Alert outside window"], ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""]],
        [0.7, 0.9, 1.35, 0.75, 0.85, 0.7, 1.0],
    )

    heading(doc, "Telemetry and Log Sources", 1)
    grid_table(
        doc,
        ["Source ID", "Product or service", "Covered assets", "Events collected", "Transport", "Retention", "Health owner", "Known gaps"],
        [["Example LS-001", "Windows Event", "Domain controllers", "Security, PowerShell", "Wazuh agent", "180 days", "SOC Platform", "None known"], ["", "", "", "", "", "", "", ""], ["", "", "", "", "", "", "", ""]],
        [0.7, 0.9, 0.9, 1.05, 0.75, 0.65, 0.8, 0.85],
    )

    heading(doc, "Security Tool Topology", 1)
    grid_table(
        doc,
        ["Platform", "Purpose", "Base URL or host", "Owners", "Connected sources", "Environment", "Notes"],
        [["Wazuh", "SIEM and endpoint monitoring", "https://wazuh.local", "SOC Platform", "Agents and network logs", "Production", "API account is read-only for validation"], ["TheHive", "Case management", "", "", "", "", ""], ["Cortex", "Analysis and response", "", "", "", "", ""], ["MISP", "Threat intelligence", "", "", "", "", ""]],
        [0.75, 1.1, 1.05, 0.85, 1.15, 0.75, 1.0],
    )

    heading(doc, "Wazuh Configuration and Rules", 1)
    form_table(
        doc,
        [
            ("Manager and indexer topology", "Nodes, clusters, versions, high availability, and index patterns."),
            ("Agent enrollment", "Enrollment method, groups, ownership, expected agent count, and stale-agent threshold."),
            ("Rule and decoder paths", "Locations and ownership of custom rules, decoders, lists, and integrations."),
            ("Retention and rollover", "Index lifecycle, hot and archive retention, backup, and restoration expectations."),
            ("Health monitoring", "How ingestion delays, agent failures, index failures, and dropped events are detected."),
        ],
    )
    grid_table(
        doc,
        ["Rule ID", "Rule name", "Version", "Log sources", "ATT&CK", "Severity", "Expected volume", "Owner and review date"],
        [["5710", "Example authentication detection", "1.0", "Windows Event", "T1110", "High", "0 to 20 per day", "SOC Detection / YYYY-MM-DD"], ["", "", "", "", "", "", "", ""], ["", "", "", "", "", "", "", ""]],
        [0.6, 1.25, 0.55, 0.85, 0.65, 0.6, 0.85, 1.0],
    )

    heading(doc, "TheHive Configuration", 1)
    form_table(
        doc,
        [
            ("Organization and roles", "Organizations, analyst roles, managers, service accounts, and access boundaries."),
            ("Case templates", "Templates, tasks, severities, tags, custom fields, and required closure dispositions."),
            ("Alert workflow", "Alert import, deduplication, ownership, merging, escalation, and closure process."),
            ("Observable handling", "Supported types, TLP/PAP defaults, tagging, retention, and duplicate behavior."),
            ("Audit and retention", "Audit retention, case retention, export, backup, and legal hold requirements."),
        ],
    )

    heading(doc, "Cortex Configuration", 1)
    grid_table(
        doc,
        ["Analyzer or responder", "Type", "Observable types", "Organization", "Timeout", "Approval needed", "Owner", "Known limitations"],
        [["Example VirusTotal analyzer", "Analyzer", "IP, domain, hash, URL", "SOC", "120 seconds", "No", "SOC Platform", "Vendor rate limit"], ["", "", "", "", "", "", "", ""], ["", "", "", "", "", "", "", ""]],
        [1.2, 0.65, 1.05, 0.7, 0.65, 0.8, 0.7, 0.95],
    )

    heading(doc, "MISP Configuration", 1)
    form_table(
        doc,
        [
            ("Organizations and sharing groups", "Record sharing boundaries and approved recipients."),
            ("Distribution defaults", "State event and attribute distribution, TLP handling, and publication approval."),
            ("Feeds and sync servers", "Names, owners, schedules, trust, filtering, and failure notification."),
            ("Taxonomies and galaxies", "Required taxonomies, ATT&CK galaxies, local tags, and update ownership."),
            ("Warning lists", "Enabled warning lists, update schedule, overrides, and review process."),
            ("Retention and decay", "Decay models, expiration, stale-intelligence review, and deletion policy."),
        ],
    )

    heading(doc, "Detection Ground Truth", 1)
    intro(doc, "Ground truth must come from a verified investigation, approved simulation, known benign test, incident report, or another independent source. An alert alone is not ground truth.")
    grid_table(
        doc,
        ["Event ID", "Date and time UTC", "Verified outcome", "Alert generated", "Rule expected", "Evidence source", "Case or ticket", "Approved by"],
        [["Example GT-001", "YYYY-MM-DD HH:MM", "Malicious", "No", "5710", "Purple-team report", "CASE-001", "Senior Analyst"], ["", "", "", "", "", "", "", ""], ["", "", "", "", "", "", "", ""]],
        [0.65, 0.9, 0.75, 0.7, 0.7, 1.0, 0.75, 0.8],
    )
    heading(doc, "Verdict Definitions", 2)
    grid_table(
        doc,
        ["Classification", "SIEM outcome", "Ground truth", "Required evidence"],
        [["True positive", "Alert generated", "Malicious", "Confirmed malicious activity linked to the alert"], ["False positive", "Alert generated", "Benign", "Confirmed authorized or benign activity"], ["False negative", "No alert", "Malicious", "Independent malicious event and the rule that should have detected it"], ["True negative", "No alert", "Benign", "Controlled benign event and confirmation that no alert was expected or emitted"], ["Needs review", "Any", "Unknown", "Insufficient or conflicting evidence"]],
        [1.05, 1.2, 1.0, 3.45],
    )

    heading(doc, "Replay and Validation Corpus", 1)
    grid_table(
        doc,
        ["Sample ID", "Benign or malicious", "Sanitized source", "Expected rules", "Expected outcome", "Scope approval", "Retention and owner"],
        [["Example RC-001", "Malicious", "Sanitized Windows events", "5710", "Alert with high severity", "TEST-001", "180 days / Detection Team"], ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""]],
        [0.7, 0.85, 1.15, 0.85, 1.1, 0.85, 1.0],
    )

    heading(doc, "Detection Change Governance", 1)
    form_table(
        doc,
        [
            ("Rule owners", "Teams or individuals responsible for each detection family."),
            ("Minimum replay evidence", "Required benign and malicious samples before shadow deployment."),
            ("Shadow period", "Required duration, event volume, and acceptance thresholds."),
            ("Canary scope", "Approved agents, hosts, groups, or tenants for limited deployment."),
            ("Production approval", "Roles authorized to approve production changes and separation-of-duty requirements."),
            ("Rollback criteria", "Precision, recall, latency, stability, or incident thresholds requiring rollback."),
        ],
    )

    # Keep the final evidence checklist and sign-off table together so reviewers
    # do not sign on a detached continuation page.
    doc.add_page_break()
    heading(doc, "Attachments and Evidence", 1)
    checklist(
        doc,
        [
            "Current network and trust-zone diagram",
            "Asset and data-owner inventory",
            "Wazuh rule and decoder exports with secrets removed",
            "TheHive case templates and closure-disposition definitions",
            "Cortex analyzer and responder inventory",
            "MISP taxonomy, warning-list, feed, and sharing configuration exports",
            "Approved replay corpus manifest and sanitization record",
            "Relevant incident, change, and exception tickets",
        ],
    )

    heading(doc, "Final Review", 1)
    grid_table(
        doc,
        ["Review area", "Reviewer name and role", "Decision", "Date", "Comments or conditions"],
        [["Architecture", "", "Approve / Reject", "", ""], ["SOC operations", "", "Approve / Reject", "", ""], ["Detection engineering", "", "Approve / Reject", "", ""], ["Security governance", "", "Approve / Reject", "", ""]],
        [1.1, 1.45, 1.05, 0.8, 2.1],
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AI SOC Operator organization input template  |  Do not include secrets").font.size = Pt(8)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(OUTPUT)
